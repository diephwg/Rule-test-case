# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Export testcases.csv to a tester-friendly Excel workbook.

Usage:
    uv run scripts/export_xlsx.py output/<FEATURE>/testcases.csv [--analysis output/<FEATURE>/analysis.md] [-o out.xlsx]

Sheets:
    Test Cases - all columns + Design_Status (computed, needs --analysis) + execution columns
                 (Actual_Result, Bug_ID, Tester, Executed_Date), dropdowns, frozen header, autofilter
    Summary    - counts by Status / Design_Status / Priority / Scenario / Test_Type / Module (live formulas)
    RTM        - requirement -> test cases (needs --analysis)
    Conditions - test condition (CO-xx) -> test cases (needs --analysis with a conditions table)
    Clauses    - source clause (<Req_ID>-c<n>) -> CO -> test cases, with oracle state (needs a clauses table)
"""
import argparse
import csv
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from validate_testcases import ENUMS, design_status, parse_analysis, split_ids  # noqa: E402

EXEC_COLUMNS = ["Actual_Result", "Bug_ID", "Tester", "Executed_Date"]
WIDTHS = {
    "TC_ID": 16, "Module": 16, "Req_ID": 18, "Title": 45, "Test_Type": 13, "Scenario": 11,
    "Technique": 10, "Priority": 9, "Preconditions": 35, "Test_Data": 30, "Steps": 50,
    "Expected_Result": 50, "Automation": 12, "Status": 11, "Notes": 30, "Design_Status": 16,
    "Actual_Result": 35, "Bug_ID": 12, "Tester": 14, "Executed_Date": 14,
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
EXEC_FILL = PatternFill("solid", fgColor="7F6000")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
PRIORITY_FILL = {"P1": "F8CBAD", "P2": "FFE699", "P3": "C6E0B4", "P4": "DDEBF7"}
PENDING_FILL = PatternFill("solid", fgColor="FCE4D6")
BLOCKED_FILL = PatternFill("solid", fgColor="F4CCCC")
MISSING_FILL = PatternFill("solid", fgColor="F8CBAD")


def style_header(ws, headers, exec_cols=()):
    for col, name in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.fill = EXEC_FILL if name in exec_cols else HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS.get(name, 18)
    ws.freeze_panes = "B2"


def write_rows(ws, rows, start=2):
    for r_idx, vals in enumerate(rows, start=start):
        for c_idx, v in enumerate(vals, start=1):
            c = ws.cell(row=r_idx, column=c_idx, value=v)
            c.alignment = WRAP
            c.border = BORDER


def add_list_validation(ws, col_idx, values, last_row):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    letter = get_column_letter(col_idx)
    dv.add(f"{letter}2:{letter}{max(last_row, 2) + 500}")
    ws.add_data_validation(dv)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("csv_path")
    ap.add_argument("--analysis")
    ap.add_argument("-o", "--output")
    args = ap.parse_args()

    src = Path(args.csv_path)
    out = Path(args.output) if args.output else src.with_suffix(".xlsx")
    with open(src, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        base_headers = [h.strip() for h in reader.fieldnames]
        rows = [{k.strip(): (v or "").strip() for k, v in r.items() if k} for r in reader]
    a = parse_analysis(args.analysis) if args.analysis else None

    extra = (["Design_Status"] if a else []) + EXEC_COLUMNS
    headers = base_headers + [c for c in extra if c not in base_headers]
    for row in rows:
        if a:
            status, reasons = design_status(row, a)
            row["Design_Status"] = status + (f" ({', '.join(reasons)})" if reasons else "")
        if not row.get("Status"):
            row["Status"] = "Not Run"

    wb = Workbook()

    # --- Test Cases
    ws = wb.active
    ws.title = "Test Cases"
    style_header(ws, headers, EXEC_COLUMNS)
    write_rows(ws, [[row.get(h, "") for h in headers] for row in rows])
    for r_idx, row in enumerate(rows, start=2):
        prio = row.get("Priority", "")
        if prio in PRIORITY_FILL:
            ws.cell(row=r_idx, column=headers.index("Priority") + 1).fill = PatternFill("solid", fgColor=PRIORITY_FILL[prio])
        if row.get("Design_Status", "").startswith("Needs"):
            ws.cell(row=r_idx, column=headers.index("Design_Status") + 1).fill = PENDING_FILL
        elif row.get("Design_Status", "").startswith("Blocked"):
            ws.cell(row=r_idx, column=headers.index("Design_Status") + 1).fill = BLOCKED_FILL
    last = len(rows) + 1
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    for c in ("Status", "Priority", "Scenario", "Test_Type", "Technique", "Automation"):
        if c in headers:
            add_list_validation(ws, headers.index(c) + 1, sorted(v for v in ENUMS[c] if v), last)

    # --- Summary (formulas so it updates while testers fill Status)
    sm = wb.create_sheet("Summary")
    sm["A1"] = f"Tổng hợp — {src.parent.name}"
    sm["A1"].font = Font(bold=True, size=14)
    row_i = 3
    rng_end = last + 500
    groups = {"Status": sorted(v for v in ENUMS["Status"] if v),
              "Design_Status": ["Ready", "Needs Clarification*", "Blocked*", "Obsolete"]}
    for c in ("Status", "Design_Status", "Priority", "Scenario", "Test_Type", "Module"):
        if c not in headers:
            continue
        letter = get_column_letter(headers.index(c) + 1)
        values = groups.get(c) or sorted({r.get(c, "") for r in rows if r.get(c)})
        sm.cell(row=row_i, column=1, value=c).font = Font(bold=True)
        sm.cell(row=row_i, column=2, value="Số lượng").font = Font(bold=True)
        row_i += 1
        for v in values:
            sm.cell(row=row_i, column=1, value=v.rstrip("*"))
            sm.cell(row=row_i, column=2, value=f"=COUNTIF('Test Cases'!{letter}2:{letter}{rng_end},\"{v}\")")
            row_i += 1
        row_i += 1
    sm.cell(row=row_i, column=1, value="Tổng").font = Font(bold=True)
    sm.cell(row=row_i, column=2, value=f"=COUNTA('Test Cases'!A2:A{rng_end})")
    sm.column_dimensions["A"].width = 24
    sm.column_dimensions["B"].width = 12

    if a:
        active = [r for r in rows if r.get("Status") != "Obsolete"]
        active_ids = {r["TC_ID"] for r in active}

        # --- RTM
        req_to_tcs = defaultdict(list)
        for row in active:
            for rid in split_ids(row.get("Req_ID")):
                req_to_tcs[rid].append(row["TC_ID"])
        rtm = wb.create_sheet("RTM")
        style_header(rtm, ["Req_ID", "Mô tả", "Số TC", "Test cases", "Trạng thái"])
        for w, letter in zip((18, 60, 8, 50, 16), "ABCDE"):
            rtm.column_dimensions[letter].width = w
        data = []
        for rid in sorted(set(a["reqs"]) | set(req_to_tcs)):
            tcs = req_to_tcs.get(rid, [])
            if rid in a["req_rejected"]:
                state = "Rejected"
            elif rid in a["req_na"]:
                state = "N/A"
            else:
                state = "Covered" if tcs else "NOT COVERED"
            if rid in a["imp"] and rid not in a["req_confirmed"] and state == "Covered":
                state = "Covered (IMP chưa xác nhận)"
            data.append([rid, a["reqs"].get(rid, "(không có trong analysis.md)"), len(tcs), ", ".join(tcs), state])
        write_rows(rtm, data)
        for i, vals in enumerate(data, start=2):
            if vals[4] == "NOT COVERED":
                rtm.cell(row=i, column=5).fill = MISSING_FILL

        # --- Conditions
        if a["conditions"]:
            cs = wb.create_sheet("Conditions")
            style_header(cs, ["Cond_ID", "Req_ID", "Điều kiện kiểm thử", "Kỹ thuật", "Test cases", "Trạng thái"])
            for w, letter in zip((10, 18, 60, 12, 40, 16), "ABCDEF"):
                cs.column_dimensions[letter].width = w
            data = []
            for cid, c in a["conditions"].items():
                live = [t for t in c["tcs"] if t in active_ids]
                state = "N/A" if c["na"] else "Pending" if c["pending"] else ("Covered" if live else "NOT COVERED")
                data.append([cid, "; ".join(c["reqs"]), c["desc"], c["tech"], ", ".join(live), state])
            write_rows(cs, data)
            for i, vals in enumerate(data, start=2):
                if vals[5] == "NOT COVERED":
                    cs.cell(row=i, column=6).fill = MISSING_FILL
                elif vals[5] == "Pending":
                    cs.cell(row=i, column=6).fill = PENDING_FILL

        # --- Clauses (source clause -> CO -> TC)
        if a["clauses"]:
            ready_ids = {r["TC_ID"] for r in active if r.get("Design_Status") == "Ready"}
            cl = wb.create_sheet("Clauses")
            style_header(cl, ["Clause_ID", "Req_ID", "Mệnh đề", "CO", "Test cases", "Trạng thái"])
            for w, letter in zip((16, 12, 60, 18, 40, 22), "ABCDEF"):
                cl.column_dimensions[letter].width = w
            data = []
            for cid, c in a["clauses"].items():
                live = sorted({t for co in c["cos"] for t in a["conditions"].get(co, {}).get("tcs", [])
                               if t in active_ids})
                if c["na"]:
                    state = "N/A"
                elif c["pending"] or not live and c["cos"] and all(
                        a["conditions"].get(co, {}).get("pending") for co in c["cos"]):
                    state = "Pending"
                elif not live:
                    state = "NOT COVERED"
                else:
                    state = "Covered" if set(live) & ready_ids else "Covered (oracle tạm)"
                data.append([cid, c["req"], c["desc"], "; ".join(c["cos"]), ", ".join(live), state])
            write_rows(cl, data)
            for i, vals in enumerate(data, start=2):
                if vals[5] == "NOT COVERED":
                    cl.cell(row=i, column=6).fill = MISSING_FILL
                elif vals[5] in ("Pending", "Covered (oracle tạm)"):
                    cl.cell(row=i, column=6).fill = PENDING_FILL

    wb.save(out)
    print(f"Đã xuất {len(rows)} test case → {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
