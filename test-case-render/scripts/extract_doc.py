# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Convert requirement documents to Markdown so an AI agent can read them.

Usage:
    uv run scripts/extract_doc.py input/<FEATURE> -o output/<FEATURE>/_extracted
    uv run scripts/extract_doc.py input/<FEATURE>/spec.pdf -o output/<FEATURE>/_extracted --render-pages

Supported: pdf, docx, odt, rtf, html, epub (pandoc) · doc, ppt, pptx, xls (via LibreOffice)
           xlsx, xlsm, csv · md, txt, feature, json, yaml · images (listed for direct viewing)
Writes <out>/<file>.md for each document and <out>/INDEX.md summarising everything.
External tools used when present: pdftotext, pdftoppm, pandoc, soffice.
"""
import argparse
import csv
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

PANDOC_EXT = {".docx", ".odt", ".rtf", ".html", ".htm", ".epub"}
OFFICE_CONVERT = {".doc": "docx", ".ppt": "pdf", ".pptx": "pdf", ".xls": "xlsx", ".odp": "pdf", ".ods": "xlsx"}
SHEET_EXT = {".xlsx", ".xlsm"}
TEXT_EXT = {".md", ".txt", ".feature", ".json", ".yaml", ".yml", ".xml"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".svg"}


def run(cmd):
    return subprocess.run(cmd, capture_output=True, text=True)


def md_table(rows):
    rows = [[("" if v is None else str(v)).replace("|", "\\|").replace("\n", "<br>") for v in r] for r in rows]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return "_(trống)_\n"
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines) + "\n"


def render_pages(src, out_dir):
    pages_dir = out_dir / f"{src.stem}_pages"
    pages_dir.mkdir(parents=True, exist_ok=True)
    run(["pdftoppm", "-png", "-r", "110", str(src), str(pages_dir / "page")])
    return pages_dir


def from_pdf(src, out_dir, force_render=False):
    if not shutil.which("pdftotext"):
        return None, "thiếu pdftotext (cài poppler-utils)"
    res = run(["pdftotext", "-layout", str(src), "-"])
    text = res.stdout
    note = ""
    scanned = len(text.strip()) < 50
    if scanned:
        note = "PDF không có lớp text (có thể là bản scan)"
    if (scanned or force_render) and shutil.which("pdftoppm"):
        pages_dir = render_pages(src, out_dir)
        note += f"{' → ' if note else ''}ảnh từng trang trong `{pages_dir.name}/` — xem trực tiếp để đối chiếu"
    pages = text.split("\f")
    body = "\n".join(f"\n<!-- trang {i} -->\n{p}" for i, p in enumerate(pages, start=1) if p.strip())
    return f"```text\n{body}\n```\n" if body else "", note


def from_pandoc(src, out_dir):
    if not shutil.which("pandoc"):
        return None, "thiếu pandoc"
    media = out_dir / f"{src.stem}_media"
    res = run(["pandoc", str(src), "-t", "gfm", "--wrap=none", f"--extract-media={media}"])
    if res.returncode != 0:
        return None, f"pandoc lỗi: {res.stderr.strip()[:200]}"
    note = ""
    if media.exists():
        imgs = [p for p in media.rglob("*") if p.suffix.lower() in IMAGE_EXT | {".emf", ".wmf"}]
        if imgs:
            note = f"có {len(imgs)} ảnh nhúng trong `{media.name}/` — hãy xem các ảnh này (mockup, sơ đồ)"
    return res.stdout, note


def from_sheet(src):
    from openpyxl import load_workbook

    wb = load_workbook(src, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        parts.append(f"## Sheet: {ws.title}\n\n{md_table(rows)}")
    return "\n".join(parts), ""


def from_csv(src):
    with open(src, encoding="utf-8-sig", newline="") as f:
        return md_table(list(csv.reader(f))), ""


def office_convert(src, target, tmp):
    if not shutil.which("soffice"):
        return None
    run(["soffice", "--headless", "--convert-to", target, "--outdir", str(tmp), str(src)])
    out = Path(tmp) / f"{src.stem}.{target}"
    return out if out.exists() else None


def convert(src, out_dir, force_render=False):
    ext = src.suffix.lower()
    if ext in IMAGE_EXT:
        return None, "ẢNH — mở và xem trực tiếp"
    with tempfile.TemporaryDirectory() as tmp:
        work = src
        if ext in OFFICE_CONVERT:
            work = office_convert(src, OFFICE_CONVERT[ext], tmp)
            if not work:
                return None, "không chuyển đổi được (thiếu LibreOffice/soffice)"
            ext = work.suffix.lower()
        if ext == ".pdf":
            return from_pdf(work, out_dir, force_render)
        if ext in PANDOC_EXT:
            return from_pandoc(work, out_dir)
        if ext in SHEET_EXT:
            return from_sheet(work)
        if ext == ".csv":
            return from_csv(work)
        if ext in TEXT_EXT:
            return work.read_text(encoding="utf-8", errors="replace"), ""
    return None, "định dạng chưa hỗ trợ — yêu cầu tester export sang PDF/DOCX"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", help="file hoặc thư mục trong input/")
    ap.add_argument("-o", "--out", required=True, help="thư mục đầu ra, vd. output/<FEATURE>/_extracted")
    ap.add_argument("--render-pages", action="store_true",
                    help="luôn render ảnh từng trang PDF/PPTX (dùng khi text bị lỗi dấu tiếng Việt, có bảng/sơ đồ phức tạp)")
    args = ap.parse_args()

    source = Path(args.source)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    files = sorted(p for p in source.rglob("*") if p.is_file()) if source.is_dir() else [source]
    files = [p for p in files if not p.name.startswith((".", "~$"))]

    index = ["# Tài liệu đã trích xuất\n", "| Nguồn | Kết quả | Ghi chú |", "|---|---|---|"]
    for src in files:
        content, note = convert(src, out_dir, args.render_pages)
        result = "-"
        if content is not None:
            dest = out_dir / f"{src.name}.md"
            dest.write_text(f"<!-- nguồn: {src} -->\n# {src.name}\n\n{content}", encoding="utf-8")
            result = f"`{dest.name}` ({len(content):,} ký tự)"
            if not content.strip():
                note = (note + "; " if note else "") + "KHÔNG trích được nội dung"
        elif src.suffix.lower() in IMAGE_EXT:
            result = f"`{src}`"
        index.append(f"| `{src}` | {result} | {note} |")
        print(f"{src} -> {result} {('· ' + note) if note else ''}")

    (out_dir / "INDEX.md").write_text("\n".join(index) + "\n", encoding="utf-8")
    print(f"\nINDEX: {out_dir / 'INDEX.md'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
